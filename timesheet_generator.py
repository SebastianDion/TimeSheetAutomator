import os
import json
import random
import calendar
from datetime import datetime, date, time
from typing import Dict, Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter

INDONESIAN_MONTHS = {
    1: "Januari",
    2: "Februari",
    3: "Maret",
    4: "April",
    5: "Mei",
    6: "Juni",
    7: "Juli",
    8: "Agustus",
    9: "September",
    10: "Oktober",
    11: "November",
    12: "Desember"
}

def load_json(filepath: str) -> dict:
    if os.path.exists(filepath):
        with open(filepath, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}

def set_box_borders(ws, start_row: int, end_row: int, start_col: int, end_col: int, border_side: Side):
    """Applies clean outer rectangular borders to a merged range of cells."""
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            top = border_side if r == start_row else None
            bottom = border_side if r == end_row else None
            left = border_side if c == start_col else None
            right = border_side if c == end_col else None
            cell.border = Border(top=top, bottom=bottom, left=left, right=right)

class TimesheetGenerator:
    def __init__(self, config_path: str = "config.json", holidays_path: str = "holidays.json"):
        self.config_path = config_path
        self.holidays_path = holidays_path
        self.config = load_json(config_path)
        if not self.config:
            fallback = os.path.join(os.path.dirname(config_path) or ".", "config.example.json")
            if os.path.exists(fallback):
                self.config = load_json(fallback)
        self.holidays_data = load_json(holidays_path)

    def get_holiday_name(self, date_str: str) -> Optional[str]:
        year = date_str.split("-")[0]
        year_holidays = self.holidays_data.get(year, {})
        if date_str in year_holidays:
            return year_holidays[date_str]
        custom_holidays = self.holidays_data.get("custom_holidays", {})
        if date_str in custom_holidays:
            return custom_holidays[date_str]
        return None

    def generate_random_working_hours(self) -> tuple[time, time]:
        wh = self.config.get("working_hours", {})
        if not wh.get("randomize_minutes", True):
            return time(8, 0), time(17, 15)

        start_min_h = wh.get("start_min_hour", 7)
        start_min_m = wh.get("start_min_minute", 55)
        start_max_h = wh.get("start_max_hour", 8)
        start_max_m = wh.get("start_max_minute", 15)

        start_total_mins = random.randint(start_min_h * 60 + start_min_m, start_max_h * 60 + start_max_m)
        start_h, start_m = divmod(start_total_mins, 60)

        end_min_h = wh.get("end_min_hour", 17)
        end_min_m = wh.get("end_min_minute", 5)
        end_max_h = wh.get("end_max_hour", 17)
        end_max_m = wh.get("end_max_minute", 30)

        end_total_mins = random.randint(end_min_h * 60 + end_min_m, end_max_h * 60 + end_max_m)
        end_h, end_m = divmod(end_total_mins, 60)

        return time(start_h, start_m), time(end_h, end_m)

    def parse_time_str(self, val: str) -> Optional[time]:
        if not val:
            return None
        parts = val.split(":")
        try:
            h = int(parts[0])
            m = int(parts[1]) if len(parts) > 1 else 0
            return time(h, m)
        except Exception:
            return None

    def generate_workbook(self, year: int, month: int, activities: Dict[int, Any]) -> Workbook:
        wb = Workbook()
        ws = wb.active
        month_name = INDONESIAN_MONTHS.get(month, calendar.month_name[month])
        ws.title = f"{month_name}"

        ws.views.sheetView[0].showGridLines = True

        # Styles
        font_regular = Font(name="Calibri", size=10, bold=False)
        font_bold = Font(name="Calibri", size=10, bold=True)

        fill_header = PatternFill(start_color="FFD6DCE4", end_color="FFD6DCE4", fill_type="solid")
        fill_holiday = PatternFill(start_color="FFBFBFBF", end_color="FFBFBFBF", fill_type="solid")

        thin_side = Side(border_style="thin", color="000000")
        border_all_thin = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        align_center_center = Alignment(horizontal="center", vertical="center")
        align_center_center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_center_bottom = Alignment(horizontal="center", vertical="bottom")

        emp = self.config.get("employee", {})
        client = self.config.get("client", {})
        proj = self.config.get("project_defaults", {})
        sups = self.config.get("supervisors", {})

        # --- 1. Top Metadata (Rows 1 to 6) ---
        ws["A1"] = "NAME of PROJECT"
        ws["A1"].font = font_regular
        ws["C1"] = f": {client.get('project_header', 'PT BANK NEGARA INDONESIA (Persero) Tbk')}"
        ws["C1"].font = font_bold
        ws.merge_cells("A1:B1")

        ws["A2"] = "UNIT/DIVISION"
        ws["A2"].font = font_regular
        ws["C2"] = f": {client.get('unit_division', 'RMS/WDL')}"
        ws["C2"].font = font_bold
        ws.merge_cells("A2:B2")
        ws.merge_cells("C2:E2")

        ws["A3"] = "NAME"
        ws["A3"].font = font_regular
        ws["C3"] = f"{emp.get('name', 'YOUR_NAME')}"
        ws["C3"].font = font_bold
        ws.merge_cells("A3:B3")
        ws.merge_cells("C3:F3")

        ws["A4"] = "MII ID"
        ws["A4"].font = font_regular
        ws["C4"] = f": {emp.get('mii_id', 'YOUR_MII_ID')}"
        ws["C4"].font = font_bold
        ws.merge_cells("A4:B4")


        ws["A5"] = "SITE"
        ws["A5"].font = font_regular
        ws["C5"] = f": {client.get('site', 'PT BANK NEGARA INDONESIA (PERSERO) Tbk.')}"
        ws["C5"].font = font_bold
        ws.merge_cells("A5:B5")
        ws.merge_cells("C5:F5")

        # Holiday Legend in Row 5: G5 has gray box, H5:I5 has ' : Holiday'
        ws["G5"].fill = fill_holiday
        ws["G5"].border = border_all_thin
        ws["H5"] = " : Holiday"
        ws["H5"].font = font_regular
        ws.merge_cells("H5:I5")

        ws["A6"] = "PERIODE"
        ws["A6"].font = font_regular
        ws["C6"] = f": {month_name}-{year}"
        ws["C6"].font = font_bold
        ws.merge_cells("A6:B6")
        ws.merge_cells("C6:D6")

        ws["G6"] = "P = Present; S = Sick;  V = Vacation; BT = Business Trip; PM = Permit; X = Not Working Anymore"
        ws["G6"].font = font_regular

        # Row heights for metadata
        ws.row_dimensions[1].height = 14.4
        for r in range(2, 9):
            ws.row_dimensions[r].height = 13.8

        # --- 2. Table Headers (Rows 7 & 8) ---
        headers_row7 = {
            1: "DATE",
            2: "WORKING HOUR",
            4: "TOTAL HOUR",
            5: "STATUS  ATTENDANCE ",
            11: "ACTIVITY / REMARK",
            12: "Project Name",
            13: "Project ID",
            14: "Aplikasi Terdampak",
            15: "AIP Fitur",
            16: "Divisi",
            17: "Departement",
            18: "Sub Departement"
        }
        for col_idx, val in headers_row7.items():
            cell = ws.cell(row=7, column=col_idx, value=val)

        headers_row8 = {
            2: "START",
            3: "END",
            5: "Present",
            6: "Sick ",
            7: "Business Trip",
            8: "Permit",
            9: "Vacation",
            10: "Not Working"
        }
        for col_idx, val in headers_row8.items():
            ws.cell(row=8, column=col_idx, value=val)

        ws.merge_cells("A7:A8")
        ws.merge_cells("B7:C7")
        ws.merge_cells("D7:D8")
        ws.merge_cells("E7:J7")
        ws.merge_cells("K7:K8")
        ws.merge_cells("L7:L8")
        ws.merge_cells("M7:M8")
        ws.merge_cells("N7:N8")
        ws.merge_cells("O7:O8")
        ws.merge_cells("P7:P8")
        ws.merge_cells("Q7:Q8")
        ws.merge_cells("R7:R8")

        for col_idx in range(1, 19):
            for row_idx in (7, 8):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = font_bold
                cell.fill = fill_header
                cell.alignment = align_center_center
                cell.border = border_all_thin

        # --- 3. Data Rows (Days 1 to N) ---
        num_days = calendar.monthrange(year, month)[1]
        start_data_row = 9
        last_working_day = 1

        for day in range(1, num_days + 1):
            curr_date = date(year, month, day)
            date_str = curr_date.strftime("%Y-%m-%d")
            row_idx = start_data_row + day - 1
            weekday = curr_date.weekday()
            is_weekend = (weekday in (5, 6))
            holiday_name = self.get_holiday_name(date_str)

            user_entry = activities.get(day, None)
            act_text = ""
            status_code = ""
            custom_start = None
            custom_end = None
            row_proj_name = proj.get("project_name", "CPR ")
            row_proj_id = proj.get("project_id", "CR260028")
            row_app = proj.get("aplikasi_terdampak", "CPR ")
            row_aip = proj.get("aip_fitur", "-")
            row_div = proj.get("divisi", "WDL")
            row_dept = proj.get("departement", "WEL")
            row_subdept = proj.get("sub_departement", "-")

            if isinstance(user_entry, str):
                act_text = user_entry
                if act_text.strip():
                    status_code = "P"
            elif isinstance(user_entry, dict):
                act_text = user_entry.get("activity", "")
                status_code = user_entry.get("status", "P" if act_text.strip() else "")
                custom_start = self.parse_time_str(user_entry.get("start"))
                custom_end = self.parse_time_str(user_entry.get("end"))
                row_proj_name = user_entry.get("project_name", row_proj_name)
                row_proj_id = user_entry.get("project_id", row_proj_id)
                row_app = user_entry.get("aplikasi_terdampak", row_app)
                row_div = user_entry.get("divisi", row_div)
                row_dept = user_entry.get("departement", row_dept)

            ws.row_dimensions[row_idx].height = 30.0

            # If weekend/holiday without activity, OR status is explicitly empty/kosong without activity
            is_empty_status = (status_code == "" or status_code is None)
            if (is_empty_status and not act_text) or ((is_weekend or holiday_name) and not act_text):
                for c in range(1, 19):
                    cell = ws.cell(row=row_idx, column=c, value=None)
                    cell.fill = fill_holiday
                    cell.border = border_all_thin
                continue

            # Standard working day (or holiday with work)
            last_working_day = day
            if not status_code:
                status_code = "P"

            # Col A: Date
            c_date = ws.cell(row=row_idx, column=1, value=datetime(year, month, day))
            c_date.number_format = "d-mmm-yy"
            c_date.font = font_bold
            c_date.alignment = align_center_center
            c_date.border = border_all_thin

            # Col B & C: Start & End Times
            if status_code in ("P", "BT"):
                if custom_start and custom_end:
                    s_time, e_time = custom_start, custom_end
                else:
                    s_time, e_time = self.generate_random_working_hours()

                cB = ws.cell(row=row_idx, column=2, value=s_time)
                cB.number_format = "h:mm"
                cB.font = font_bold
                cB.alignment = align_center_center
                cB.border = border_all_thin

                cC = ws.cell(row=row_idx, column=3, value=e_time)
                cC.number_format = "h:mm"
                cC.font = font_bold
                cC.alignment = align_center_center
                cC.border = border_all_thin

                # Col D: Total Hour formula =C{row}-B{row}
                cD = ws.cell(row=row_idx, column=4, value=f"=C{row_idx}-B{row_idx}")
                cD.number_format = "h:mm"
                cD.font = font_bold
                cD.alignment = align_center_center
                cD.border = border_all_thin
            else:
                # e.g. Sick, Vacation
                for c in (2, 3, 4):
                    cell = ws.cell(row=row_idx, column=c, value=None)
                    cell.border = border_all_thin

            # Col E..J: Attendance Status matrix
            st_upper = status_code.upper()
            status_cols = {
                "P": 5,
                "S": 6,
                "BT": 7,
                "PM": 8,
                "V": 9,
                "X": 10
            }
            target_col = status_cols.get(st_upper, 5 if act_text else None)
            for c in range(5, 11):
                val = st_upper if c == target_col else None
                cell = ws.cell(row=row_idx, column=c, value=val)
                cell.font = font_bold
                cell.alignment = align_center_center
                cell.border = border_all_thin

            # Col K: Activity / Remark
            cK = ws.cell(row=row_idx, column=11, value=act_text)
            cK.font = font_bold
            cK.alignment = align_center_center_wrap
            cK.border = border_all_thin

            # Col L..R: Project info
            proj_vals = [row_proj_name, row_proj_id, row_app, row_aip, row_div, row_dept, row_subdept]
            for i, p_val in enumerate(proj_vals):
                cell = ws.cell(row=row_idx, column=12 + i, value=p_val)
                cell.font = font_bold
                cell.alignment = align_center_center
                cell.border = border_all_thin

        # --- 4. Separator & Total Summary Rows ---
        last_data_row = start_data_row + num_days - 1
        sep_row = last_data_row + 1
        ws.row_dimensions[sep_row].height = 13.8

        # Separator row has borders around columns
        for c in range(1, 11):
            ws.cell(row=sep_row, column=c).border = border_all_thin
        for c in range(12, 19):
            ws.cell(row=sep_row, column=c).border = Border(top=thin_side, left=thin_side, right=thin_side)

        total_row = sep_row + 1
        ws.row_dimensions[total_row].height = 13.8

        # Merge A:D on total row
        ws.merge_cells(f"A{total_row}:D{total_row}")
        for c in range(1, 5):
            ws.cell(row=total_row, column=c).border = border_all_thin

        # Attendance counts
        ws.cell(row=total_row, column=5, value=f'=COUNTIF(E{start_data_row}:E{sep_row},"P")')
        ws.cell(row=total_row, column=6, value=f'=COUNTIF(F{start_data_row}:F{sep_row},"S")')
        ws.cell(row=total_row, column=7, value=f'=COUNTIF(G{start_data_row}:G{sep_row},"BT")')
        ws.cell(row=total_row, column=8, value=f'=COUNTIF(H{start_data_row}:H{sep_row},"PM")')
        ws.cell(row=total_row, column=9, value=f'=COUNTIF(I{start_data_row}:I{sep_row},"V")')
        ws.cell(row=total_row, column=10, value=f'=COUNTIF(J{start_data_row}:J{sep_row},"x")')

        for c in range(5, 11):
            cell = ws.cell(row=total_row, column=c)
            cell.font = font_bold
            cell.alignment = align_center_center
            cell.border = border_all_thin

        for c in range(11, 19):
            ws.cell(row=total_row, column=c).border = Border(top=thin_side)

        # --- 5. Signatures Block ---
        sign_header_row = total_row + 1
        ws.row_dimensions[sign_header_row].height = 13.8

        ws.cell(row=sign_header_row, column=1, value="TTD PEGAWAI,")
        ws.merge_cells(f"A{sign_header_row}:C{sign_header_row}")
        set_box_borders(ws, sign_header_row, sign_header_row, 1, 3, thin_side)
        ws.cell(row=sign_header_row, column=1).alignment = align_center_center
        ws.cell(row=sign_header_row, column=1).font = font_bold

        ws.cell(row=sign_header_row, column=4, value="DI PERIKSA OLEH,")
        ws.merge_cells(f"D{sign_header_row}:F{sign_header_row}")
        set_box_borders(ws, sign_header_row, sign_header_row, 4, 6, thin_side)
        ws.cell(row=sign_header_row, column=4).alignment = align_center_center
        ws.cell(row=sign_header_row, column=4).font = font_bold

        ws.cell(row=sign_header_row, column=7, value="DISETUJUI OLEH,")
        ws.merge_cells(f"G{sign_header_row}:J{sign_header_row}")
        set_box_borders(ws, sign_header_row, sign_header_row, 7, 10, thin_side)
        ws.cell(row=sign_header_row, column=7).alignment = align_center_center
        ws.cell(row=sign_header_row, column=7).font = font_bold

        # Signatures body (empty space + name at bottom)
        sign_start = sign_header_row + 1
        sign_end = sign_start + 3
        for r in range(sign_start, sign_end + 1):
            ws.row_dimensions[r].height = 13.8

        employee_sign = emp.get("signature_name", emp.get("name", "YOUR_FULL_NAME"))
        ws.cell(row=sign_start, column=1, value=f" ({employee_sign})")
        ws.merge_cells(f"A{sign_start}:C{sign_end}")
        set_box_borders(ws, sign_start, sign_end, 1, 3, thin_side)
        ws.cell(row=sign_start, column=1).alignment = align_center_bottom
        ws.cell(row=sign_start, column=1).font = font_bold

        ws.cell(row=sign_start, column=4, value=f"( {sups.get('checked_by', 'NAMA_PEMERIKSA')} )")
        ws.merge_cells(f"D{sign_start}:F{sign_end}")
        set_box_borders(ws, sign_start, sign_end, 4, 6, thin_side)
        ws.cell(row=sign_start, column=4).alignment = align_center_bottom
        ws.cell(row=sign_start, column=4).font = font_bold

        ws.cell(row=sign_start, column=7, value=f"( {sups.get('approved_by', 'NAMA_PENYETUJU')} )")
        ws.merge_cells(f"G{sign_start}:J{sign_end}")
        set_box_borders(ws, sign_start, sign_end, 7, 10, thin_side)
        ws.cell(row=sign_start, column=7).alignment = align_center_bottom
        ws.cell(row=sign_start, column=7).font = font_bold

        # Signatures Date row
        sign_date_row = sign_end + 1
        ws.row_dimensions[sign_date_row].height = 13.8

        date_str = f"DATE : {num_days:02d}/{month:02d}/{year}"
        ws.cell(row=sign_date_row, column=1, value=date_str)
        ws.merge_cells(f"A{sign_date_row}:C{sign_date_row}")
        set_box_borders(ws, sign_date_row, sign_date_row, 1, 3, thin_side)
        ws.cell(row=sign_date_row, column=1).alignment = align_center_center
        ws.cell(row=sign_date_row, column=1).font = font_regular

        ws.cell(row=sign_date_row, column=4, value="DATE : ")
        ws.merge_cells(f"D{sign_date_row}:F{sign_date_row}")
        set_box_borders(ws, sign_date_row, sign_date_row, 4, 6, thin_side)
        ws.cell(row=sign_date_row, column=4).alignment = align_center_center
        ws.cell(row=sign_date_row, column=4).font = font_regular

        ws.cell(row=sign_date_row, column=7, value="DATE : ")
        ws.merge_cells(f"G{sign_date_row}:J{sign_date_row}")
        set_box_borders(ws, sign_date_row, sign_date_row, 7, 10, thin_side)
        ws.cell(row=sign_date_row, column=7).alignment = align_center_center
        ws.cell(row=sign_date_row, column=7).font = font_regular

        # --- 6. Exact Column Widths ---
        column_widths = {
            "A": 12.66,
            "B": 13.0,
            "C": 20.0,
            "D": 20.33,
            "E": 12.66,
            "F": 13.0,
            "G": 13.0,
            "H": 13.0,
            "I": 13.0,
            "J": 13.0,
            "K": 67.44,
            "L": 45.0,
            "M": 12.66,
            "N": 13.0,
            "O": 13.0,
            "P": 13.0,
            "Q": 13.0,
            "R": 13.0
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        # --- 7. Embed Logo MII ---
        logo_path = os.path.join(os.path.dirname(__file__), "assets", "mii_logo.png")
        if os.path.exists(logo_path):
            try:
                img = Image(logo_path)
                img.width = 133
                img.height = 96
                ws.add_image(img, "M1")
            except Exception as e:
                print(f"Warning: Could not add image: {e}")

        return wb

def generate_timesheet(config_path: str, holidays_path: str, year: int, month: int, activities: Dict[int, Any], output_path: str) -> str:
    gen = TimesheetGenerator(config_path, holidays_path)
    wb = gen.generate_workbook(year, month, activities)
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return output_path
