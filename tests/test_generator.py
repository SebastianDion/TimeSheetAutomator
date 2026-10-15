import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from timesheet_generator import TimesheetGenerator, generate_timesheet

def test_september_2026_generation():
    config_path = "config.json"
    holidays_path = "holidays.json"
    output_path = "output/test_september_2026.xlsx"
    
    activities = {
        1: {"activity": "Mengerjakan fixing pada fitur simulasi BE", "start": "08:02", "end": "18:00"},
        2: {"activity": "Mengerjakan fixing pada fitur simulasi BE", "start": "08:15", "end": "17:15"},
        3: {"activity": "Upgrade Tech Stack pada Master Jabatan", "start": "08:11", "end": "17:10"},
        4: {"activity": "Menambahkan unittest pada Master Jabatan", "start": "07:59", "end": "17:37"},
    }

    res_path = generate_timesheet(config_path, holidays_path, 2026, 9, activities, output_path)
    assert os.path.exists(res_path), f"File {res_path} was not created!"
    
    wb = openpyxl.load_workbook(res_path)
    ws = wb.active
    
    assert ws["A1"].value == "NAME of PROJECT"
    assert ws["A3"].value == "NAME"
    assert ws["C3"].value is not None
    assert ws["C4"].value.startswith(":")
    assert ws["C6"].value == ": September-2026"
    assert ws["A7"].value == "DATE"
    assert ws["E8"].value == "Present"
    
    # Check formula in D9
    assert ws["D9"].value == "=C9-B9"
    assert ws["E9"].value == "P"
    
    # Check weekend row 13 (Sep 5) is solid gray and empty
    assert ws["A13"].value is None
    assert ws["A13"].fill.start_color.rgb == "FFBFBFBF"
    
    # Check total row 40
    assert ws["E40"].value == '=COUNTIF(E9:E39,"P")'
    
    # Check signature boxes
    assert ws["A41"].value == "TTD PEGAWAI,"
    assert ws["A42"].value is not None
    assert ws["D42"].value is not None
    assert ws["G42"].value is not None
    assert ws["A46"].value == "DATE : 30/09/2026"

    
    # Check logo is embedded
    assert len(ws._images) == 1
    
    print("SUCCESS: All assertions passed for updated generator!")

if __name__ == "__main__":
    test_september_2026_generation()
